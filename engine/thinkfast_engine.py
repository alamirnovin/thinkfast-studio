"""Private local decision engine used by ThinkFast Studio's desktop installer."""
from __future__ import annotations

import base64
import io
import re
import sys
from pathlib import Path
from threading import Lock, Thread
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from laya import Router

app = FastAPI(title="ThinkFast Studio Engine")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://tauri.localhost", "https://tauri.localhost", "tauri://localhost", "http://localhost:1420"],
    allow_methods=["*"],
    allow_headers=["*"],
)
router: Router | None = None
lock = Lock()
install_lock = Lock()
install_status = {
    "state": "starting",
    "percent": 0,
    "message": "Starting the included decision engine…",
    "error": None,
}

class AnalyzeRequest(BaseModel):
    records: list[dict]


class ExtractDocumentRequest(BaseModel):
    filename: str
    content: str


def legacy_word_text(content: bytes) -> str:
    """Recover readable text from legacy .doc files without requiring Word on the user's computer."""
    try:
        import olefile

        document = olefile.OleFileIO(io.BytesIO(content)).openstream("WordDocument").read()
    except Exception as exc:
        raise ValueError("This .doc file is not a readable Microsoft Word document.") from exc
    candidates = [document.decode("utf-16-le", errors="ignore"), document.decode("cp1252", errors="ignore")]
    text = max(candidates, key=lambda value: sum(character.isprintable() for character in value))
    text = re.sub(r"[\x00-\x08\x0b-\x1f]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) < 20:
        raise ValueError("ThinkFast Studio could not recover readable text from this legacy .doc file. Save it as .docx and upload it again.")
    return text


def document_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        try:
            from docx import Document

            document = Document(io.BytesIO(content))
            sections = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                sections.extend(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()) for row in table.rows)
            text = "\n".join(section for section in sections if section)
        except Exception as exc:
            raise ValueError("This .docx file could not be read.") from exc
    elif suffix == ".doc":
        text = legacy_word_text(content)
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            text = "\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
        except Exception as exc:
            raise ValueError("This PDF could not be read. It may be damaged or needs an OCR copy.") from exc
        if len(re.sub(r"\s+", "", text)) < 20:
            raise ValueError("No readable text was found in this PDF. It may be a scanned PDF; use OCR to make a searchable PDF, then upload it again.")
    else:
        raise ValueError("Only Word documents and PDFs can be extracted here.")
    if not text.strip():
        raise ValueError("This Word document does not contain readable body text.")
    return text


def set_install_status(state: str, percent: int, message: str, error: str | None = None):
    with install_lock:
        install_status.update({
            "state": state,
            "percent": max(0, min(100, int(percent))),
            "message": message,
            "error": error,
        })


def bundled_model_directory() -> Path:
    bundle_root = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return bundle_root / "models" / "english"


def load_included_model():
    global router
    try:
        model_directory = bundled_model_directory()
        if not model_directory.is_dir():
            raise FileNotFoundError("The included decision model was not found in this installation.")
        set_install_status("starting", 25, "Loading the included decision engine…")
        with lock:
            loaded_router = Router(models={"english": str(model_directory)}, default="english", max_loaded=1)
            warmup = {"ready": {"type": "noul", "instructions": "Is this text present?"}}
            loaded_router.predict("ThinkFast Studio is ready.", warmup, model="english")
            router = loaded_router
        set_install_status("ready", 100, "Included decision engine ready.")
    except Exception as exc:
        router = None
        set_install_status("failed", 0, "The included decision engine could not start.", str(exc))

def questions_for_local_model(questions: list[dict]) -> dict:
    built = {}
    for index, question in enumerate(questions):
        kind = question.get("type", "choice")
        key = f"question_{index + 1}"
        if kind == "yesno":
            built[key] = {"type": "noul", "instructions": question["text"]}
        elif kind == "score":
            built[key] = {"type": "score", "instructions": question["text"], "criteria": question.get("options") or ["low", "high"]}
        else:
            choices = question.get("options") or ["Yes", "No"]
            built[key] = {"type": "choice", "instructions": question["text"], "criteria": {choice: choice for choice in choices}}
    return built

@app.get("/health")
def health():
    with install_lock:
        status = dict(install_status)
    return {"ready": router is not None, "install": status}


@app.post("/extract-document")
def extract_document(request: ExtractDocumentRequest):
    try:
        content = base64.b64decode(request.content, validate=True)
        return {"text": document_text(request.filename, content)}
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

@app.post("/analyze")
def analyze(request: AnalyzeRequest):
    if router is None:
        raise HTTPException(status_code=409, detail="The included decision engine is still starting.")
    results = []
    for record in request.records:
        questions = record.get("questions") or []
        if not questions:
            raise HTTPException(status_code=422, detail="Choose at least one question before starting an analysis.")
        output = router.predict(record["text"], questions_for_local_model(questions))
        answers = []
        for index, question in enumerate(questions):
            raw_answer = output["answers"].get(f"question_{index + 1}", {})
            answer = raw_answer.get("choice", raw_answer.get("noul", raw_answer.get("score", "Needs review")))
            confidence = round(float(raw_answer.get("confidence", raw_answer.get("probability", 0.5))) * 100)
            answer_entry = {
            "question_id": str(question.get("id", f"question_{index + 1}")),
            "question": str(question.get("text", f"Question {index + 1}")),
            "type": kind,
            "answer": str(answer),
            "confidence": confidence,
        }
        if kind == "score":
            levels = [str(level) for level in question.get("options") or []]
            answer_entry["scale_levels"] = levels
            score_position = None
            try:
                score_position = float(raw_answer.get("score"))
            except (TypeError, ValueError):
                normalized_answer = str(answer).strip().casefold()
                for position, level in enumerate(levels, start=1):
                    if normalized_answer == level.strip().casefold():
                        score_position = float(position)
                        break
            answer_entry["score_position"] = score_position
        answers.append(answer_entry)
        first = answers[0]
        overall_confidence = min(answer["confidence"] for answer in answers)
        results.append({"title": record["title"], "text": record["text"], "answer": first["answer"], "confidence": overall_confidence, "answers": answers, "status": "ready" if overall_confidence >= 75 else "review"})
    return results

if __name__ == "__main__":
    import uvicorn

    Thread(target=load_included_model, daemon=True).start()
    uvicorn.run(app, host="127.0.0.1", port=8765, log_level="warning")
